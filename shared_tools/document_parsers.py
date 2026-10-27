"""Bounded extraction from uploaded bytes; never opens user-provided paths."""

import csv
from io import BytesIO, StringIO
from pathlib import PurePosixPath
import threading
import warnings
import zipfile

from defusedxml import ElementTree
from openpyxl import load_workbook
from PIL import Image
import pymupdf

MAX_FILE_BYTES = 20 * 1024 * 1024
TEXT_LIMIT = 20000
_PDF_LOCK = threading.Lock()


class DocumentError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


def _decode(data: bytes) -> str:
    encodings = ["utf-16"] if data.startswith((b"\xff\xfe", b"\xfe\xff")) else ["utf-8-sig", "gb18030"]
    for encoding in encodings:
        try:
            return data.decode(encoding)
        except UnicodeError:
            continue
    raise DocumentError("invalid_encoding", "无法识别文字编码，请导出为 UTF-8 文本后重试。")


def _zip_check(data: bytes) -> set[str]:
    with zipfile.ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(item.file_size for item in entries) > 50 * 1024 * 1024:
            raise DocumentError("archive_limit", "文档解压后过大或包含过多条目，请拆分后上传。")
        return set(archive.namelist())


def _word(data: bytes) -> dict:
    with zipfile.ZipFile(BytesIO(data)) as archive:
        body = ElementTree.fromstring(archive.read("word/document.xml"))
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    container = body.find(f"{ns}body")
    sections = []
    if container is not None:
        for block in container:
            if block.tag == f"{ns}tbl":
                lines = []
                for row in block.findall(f"{ns}tr"):
                    lines.append("\t".join("".join(cell.itertext()) for cell in row.findall(f"{ns}tc")))
                text = "\n".join(lines)
                unit = "table"
            else:
                text = "".join(block.itertext())
                unit = "paragraph"
            if text.strip():
                sections.append({"unit": unit, "index": len(sections) + 1, "text": text})
    return {"format": "docx", "sections": sections, "images": [], "warnings": ["Word 提取正文和表格文本，不解析批注、嵌入图片和复杂版式。"]}


def _excel(data: bytes) -> dict:
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    sections = []
    messages = ["Excel 提取单元格原始值，不复现显示格式、合并单元格或图表。"]
    truncated = len(workbook.worksheets) > 5
    try:
        for sheet in workbook.worksheets[:5]:
            lines = []
            for row in sheet.iter_rows(min_row=1, max_row=200, max_col=40):
                values = []
                for cell in row:
                    value = cell.value
                    if cell.data_type == "f":
                        messages.append("表格含公式：保留公式原文，未计算其结果。")
                    values.append("" if value is None else str(value))
                while values and values[-1] == "":
                    values.pop()
                if values:
                    lines.append("\t".join(values))
            if (sheet.max_row or 0) > 200 or (sheet.max_column or 0) > 40:
                truncated = True
            sections.append({"unit": "sheet", "name": sheet.title, "text": "\n".join(lines)})
    finally:
        workbook.close()
    if truncated:
        messages.append("表格最多读取前 5 个工作表、每表 200 行和 40 列；其余内容未读取。")
    return {"format": "xlsx", "sections": sections, "images": [], "warnings": sorted(set(messages)), "truncated": truncated}


def _image(data: bytes) -> tuple[bytes, str]:
    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(BytesIO(data)) as image:
            if image.format not in {"JPEG", "PNG", "GIF", "WEBP"}:
                raise DocumentError("unsupported_format", "图片仅支持 JPEG、PNG、GIF、WebP。")
            if image.width * image.height > 24_000_000:
                raise DocumentError("image_limit", "图片像素过大，请缩小或裁剪后上传。")
            image.seek(0)
            rgb = image.convert("RGB")
            rgb.thumbnail((1800, 1800))
            output = BytesIO()
            rgb.save(output, format="PNG")
            return output.getvalue(), "image/png"


def _pdf(data: bytes, start_page: int, max_pages: int) -> dict:
    sections, images = [], []
    with _PDF_LOCK, pymupdf.open(stream=data, filetype="pdf") as document:
        if document.needs_pass:
            raise DocumentError("password_required", "PDF 有密码保护，请在本机解锁后上传，不要在聊天中发送密码。")
        total = len(document)
        if start_page > total:
            raise DocumentError("invalid_page", "起始页超出了 PDF 总页数。")
        end = min(total, start_page - 1 + max_pages)
        for index in range(start_page - 1, end):
            page = document[index]
            text = page.get_text("text", sort=True).strip()
            if len(text) >= 40 or (text and not page.get_images()):
                sections.append({"unit": "page", "index": index + 1, "text": text, "method": "local_pdf_text"})
            else:
                # Bound raster size before rendering untrusted page dimensions.
                longest = max(page.rect.width, page.rect.height)
                if longest <= 0:
                    raise DocumentError("invalid_pdf", "PDF 页面尺寸无效。")
                scale = min(2.0, 1800 / longest)
                pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
                images.append((index + 1, pixmap.tobytes("png"), "image/png"))
    return {"format": "pdf", "total_pages": total, "selected_pages": [start_page, end], "sections": sections, "images": images, "truncated": start_page != 1 or end < total, "warnings": ["PDF 文字提取不保证复杂表格的行列关系；金额需结合原页核对。"]}


def extract_document(data: bytes, filename: str, mime_type: str, start_page: int, max_pages: int) -> dict:
    if not data:
        raise DocumentError("empty_file", "上传文件为空。")
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError("file_limit", "单个文件最多支持 20 MiB，请拆分后上传。")
    suffix = PurePosixPath(filename).suffix.lower()
    try:
        if b"%PDF-" in data[:1024]:
            return _pdf(data, start_page, max_pages)
        if data.startswith(b"PK"):
            entries = _zip_check(data)
            if "word/document.xml" in entries:
                return _word(data)
            if "xl/workbook.xml" in entries:
                return _excel(data)
            raise DocumentError("unsupported_format", "该 ZIP/Office 类型暂不支持，当前支持 DOCX 和 XLSX。")
        if mime_type.startswith("image/") or suffix in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
            image, mime = _image(data)
            return {"format": "image", "sections": [], "images": [(1, image, mime)], "warnings": ["图片只识别首帧；清晰度或版式可能导致文字识别误差。"]}
        if suffix in {".txt", ".md", ".csv", ".json", ".xml", ".log"} or mime_type.startswith("text/"):
            # Text size is bounded separately before parsing or passing it to a model.
            text = _decode(data)
            shortened = len(text) > 200000
            text = text[:200000]
            if suffix == ".csv":
                try:
                    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                reader = csv.reader(StringIO(text), dialect)
                lines = []
                for index, row in enumerate(reader):
                    if index >= 200:
                        shortened = True
                        break
                    shortened = shortened or len(row) > 40
                    lines.append("\t".join(row[:40]))
                text = "\n".join(lines)
            return {"format": "csv" if suffix == ".csv" else "text", "sections": [{"unit": "text", "index": 1, "text": text}], "images": [], "truncated": shortened, "warnings": ["文本/表格有读取上限，truncated=true 表示内容不完整。"] if shortened else []}
    except DocumentError:
        raise
    except Exception:
        raise DocumentError("invalid_document", "文件损坏、格式不符或无法安全解析，请重新导出后上传。") from None
    raise DocumentError("unsupported_format", "暂支持 PDF、DOCX、XLSX、CSV、TXT、Markdown 和常见图片；不支持旧版 DOC/XLS 或压缩包。")
